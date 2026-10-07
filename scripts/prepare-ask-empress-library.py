"""Extract the supplied ZIP into deduplicated, source-labelled search passages.

Read-only for the source files. Requires openpyxl and pypdf (bundled Python).
Output stays in ignored artifacts/; source documents are not published to Git.
"""
import argparse
import csv
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET

import openpyxl
from pypdf import PdfReader

MAX_CHARS = 2200
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def clean(value):
    return re.sub(r"[\t \u00a0]+", " ", unicodedata.normalize("NFKC", str(value or ""))).strip()


def fingerprint(text):
    return hashlib.sha256(re.sub(r"\s+", " ", clean(text)).lower().encode()).hexdigest()


def source_url(text):
    urls = re.findall(r"https?://[^\s<>\"\)]+", text)
    # Prefer the paper/guideline link to a social-media tracking link.
    for url in urls:
        if any(host in url for host in ["pubmed.ncbi.nlm.nih.gov", "doi.org", "menopause.org", "mayoclinic.org", "nhs.uk", "nih.gov"]):
            return url.rstrip(".,;")
    return ""


def split_text(text, limit=MAX_CHARS):
    """Keep sentences together; never truncate an answer or PDF page."""
    parts = re.split(r"(?<=[.!?])\s+|\n+", text)
    result, current = [], ""
    for part in parts:
        part = clean(part)
        if not part:
            continue
        while len(part) > limit:
            cut = part.rfind(" ", 0, limit)
            if cut < limit // 2:
                cut = limit
            if current:
                result.append(current)
                current = ""
            result.append(part[:cut].strip())
            part = part[cut:].strip()
        if len(current) + len(part) + 1 > limit:
            result.append(current)
            current = ""
        current = (current + " " + part).strip()
    if current:
        result.append(current)
    return result


def extract_archive(archive_path, output):
    records, seen = [], set()
    report = {"archive": archive_path.name, "files": [], "duplicate_passages": 0, "empty_answers": 0}

    def add(text, filename, kind, locator, title=None, category="", question="", url=""):
        text = clean(text)
        if len(text) < 35:
            return 0
        key = fingerprint(question + "\n" + text)
        if key in seen:
            report["duplicate_passages"] += 1
            return 0
        seen.add(key)
        count = 0
        title = title or PurePosixPath(filename).stem.replace("_", " ")
        prefix = "Question: " + question[:700] + "\n" if question else ""
        for i, part in enumerate(split_text(text, MAX_CHARS - len(prefix))):
            content = prefix + part
            content_hash = fingerprint(content)
            metadata = {
                "source_title": title[:250], "source_file": filename,
                "source_type": kind, "source_locator": locator,
                "category": category[:200], "content_hash": content_hash,
                "archive": archive_path.name, "import_version": "2026-10-07-v1",
            }
            if url:
                metadata["source_url"] = url
            if question:
                metadata["question"] = question[:700]
            records.append({"_id": "empress-library-" + content_hash[:32], "content": content, "metadata": metadata})
            count += 1
        return count

    def table(rows, filename, sheet=""):
        count, recognised = 0, False
        rows = list(rows)
        for header_index, header in enumerate(rows[:12]):
            labels = [clean(v).lower() for v in header]
            if "question" in labels and "answer" in labels:
                recognised = True
                qi, ai = labels.index("question"), labels.index("answer")
                ci = labels.index("category") if "category" in labels else None
                kind = "research_summary" if filename.startswith(("14 New Research Papers/", "Addition Q&A documents/")) else "educational_qa"
                for row_index, row in enumerate(rows[header_index + 1:], header_index + 2):
                    question = clean(row[qi]) if qi < len(row) else ""
                    answer = clean(row[ai]) if ai < len(row) else ""
                    if not question or not answer:
                        report["empty_answers"] += 1
                        continue
                    category = clean(row[ci]) if ci is not None and ci < len(row) else ""
                    count += add(answer, filename, kind, f"{sheet + ', ' if sheet else ''}row {row_index}",
                                 category=category, question=question, url=source_url(answer))
                break
            if "name" in labels and "state" in labels and "category" in labels:
                recognised = True
                for row_index, row in enumerate(rows[header_index + 1:], header_index + 2):
                    fields = dict(zip(labels, [clean(v) for v in row]))
                    if not fields.get("name"):
                        continue
                    # Professional public listing fields only. Do not index emails,
                    # LinkedIn profiles, workbook notes, or unrelated contact data.
                    keys = ["name", "qualification", "category", "state", "zip", "address", "phone", "website"]
                    text = ". ".join(f"{k.title()}: {fields[k]}" for k in keys if fields.get(k))
                    count += add(text, filename, "provider_directory", f"{sheet}, row {row_index}",
                                 title="Menopause provider directory", url=fields.get("website", ""))
                break
        return count, recognised

    with zipfile.ZipFile(archive_path) as archive:
        files = [i for i in archive.infolist() if not i.is_dir()]
        # Specific Q&A files before the combined workbook; CSV before XLSX copies.
        files.sort(key=lambda i: ("Combined_QA" in i.filename, i.filename.endswith(".xlsx"), i.filename))
        for item in files:
            filename = item.filename
            entry = {"file": filename, "bytes": item.file_size, "passages": 0}
            if item.file_size > 100_000_000 or ".." in PurePosixPath(filename).parts or filename.startswith("/"):
                entry["status"] = "excluded_unsafe_archive_entry"
                report["files"].append(entry)
                continue
            try:
                data = archive.read(item)
                suffix = PurePosixPath(filename).suffix.lower()
                if "_Metadata.csv" in filename:
                    entry["status"] = "metadata_only"
                elif suffix == ".csv":
                    entry["passages"], found = table(csv.reader(io.StringIO(data.decode("utf-8-sig"))), filename)
                    entry["status"] = "read" if found else "no_supported_table"
                elif suffix == ".xlsx":
                    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                    found = False
                    for sheet in workbook:
                        n, recognised = table(sheet.iter_rows(values_only=True), filename, sheet.title)
                        if not recognised and filename in {"Blogs/longevity clinic solution map.xlsx", "Q&A - Chatbot/Supplements for menopausal women.xlsx"}:
                            # These two sheets are reference tables rather than Q&A.
                            kind = "business_reference" if filename.startswith("Blogs/") else "educational_qa"
                            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                                cells = [clean(v) for v in row if v is not None]
                                n += add(" | ".join(cells), filename, kind, f"{sheet.title}, row {row_number}")
                            recognised = True
                        entry["passages"] += n
                        found = found or recognised
                    workbook.close()
                    entry["status"] = "read" if found else "no_supported_table"
                elif suffix == ".docx":
                    with zipfile.ZipFile(io.BytesIO(data)) as doc:
                        root = ET.fromstring(doc.read("word/document.xml"))
                        paragraphs = ["".join(p.itertext()).strip() for p in root.iter(W + "p")]
                        text = "\n".join(p for p in paragraphs if p)
                        links = []
                        if "word/_rels/document.xml.rels" in doc.namelist():
                            links = [e.attrib.get("Target", "") for e in ET.fromstring(doc.read("word/_rels/document.xml.rels")) if e.attrib.get("TargetMode") == "External"]
                    kind = "source_commentary" if filename.startswith("Citations/") else "editorial"
                    entry["passages"] = add(text, filename, kind, "document", url=source_url(text + " " + " ".join(links)))
                    entry["status"] = "read"
                elif suffix == ".pdf":
                    reader = PdfReader(io.BytesIO(data))
                    kind = ("brand_faq" if "FAQ" in filename else "clinical_framework" if "Framework" in filename
                            else "reference_compilation" if filename == "Empress_mvp.pdf"
                            else "business_reference" if filename.startswith("Blogs/") else "research_pdf")
                    entry["pages"] = len(reader.pages)
                    for page_index, page in enumerate(reader.pages, 1):
                        text = page.extract_text() or ""
                        entry["passages"] += add(text, filename, kind, f"page {page_index}", url=source_url(text))
                    entry["status"] = "read" if entry["passages"] else "no_extractable_text"
                else:
                    entry["status"] = "link_only_or_unsupported"
            except Exception as error:
                entry["status"] = "error"
                entry["error"] = str(error)[:250]
            report["files"].append(entry)
            if len(report["files"]) % 100 == 0:
                print(f"Read {len(report['files'])}/{len(files)} files; {len(records)} passages", flush=True)

    # Content-derived IDs are unique across sources and remain stable on reruns.
    by_id = {}
    for record in records:
        by_id.setdefault(record["_id"], record)
    report["chunk_duplicates"] = len(records) - len(by_id)
    records = list(by_id.values())
    report["total_files"] = len(report["files"])
    report["total_passages"] = len(records)
    report["source_types"] = dict(Counter(r["metadata"]["source_type"] for r in records))
    report["file_statuses"] = dict(Counter(f["status"] for f in report["files"]))
    output.mkdir(parents=True, exist_ok=True)
    (output / "corpus.json").write_text(json.dumps(records, ensure_ascii=False, indent=2))
    (output / "extraction-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps({k: v for k, v in report.items() if k != "files"}, indent=2))
    if report["file_statuses"].get("error"):
        raise SystemExit("Extraction errors require review; see extraction-report.json")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("archive", type=Path)
    parser.add_argument("--output", type=Path, default=Path("artifacts/ask-empress-library"))
    args = parser.parse_args()
    extract_archive(args.archive, args.output)
